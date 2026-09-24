#!/usr/bin/env python3
"""
Extract HR emails + phone numbers + company + job profile from phone screenshots -> CSV + Excel.
One row per company: all emails / phones of the same company are joined with commas.

Examples (OpenRouter):
    python screenshot-to-contacts.py --provider openai --model anthropic/claude-haiku-4.5 --test
    python screenshot-to-contacts.py --provider openai --model anthropic/claude-haiku-4.5 --input hr_shots --limit 3
    python screenshot-to-contacts.py --provider openai --model anthropic/claude-haiku-4.5 --input hr_shots
.env for provider "openai" (OpenRouter):
    OPENAI_BASE_URL=https://openrouter.ai/api/v1
    OPENAI_API_KEY=your-openrouter-key

.env for provider "anthropic":
    ANTHROPIC_API_KEY=your-key

Needs: pip install anthropic openai pillow openpyxl python-dotenv
"""
import argparse
import base64
import csv
import io
import json
import os
import re
import sys
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from PIL import Image

EXTS = {".png", ".jpg", ".jpeg", ".webp"}
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}$")
LEGAL_RE = re.compile(r"\b(pvt|private|ltd|limited|llp|inc|corp|corporation)\b")
FIELDS = ["company", "profile", "emails", "phones", "hr_name", "source_file"]
DEFAULT_MODELS = {"anthropic": "claude-haiku-4-5-20251001", "openai": "anthropic/claude-haiku-4.5"}
LIST_SEP = ", "   # emails and phones
TEXT_SEP = " | "  # profile, hr_name, source_file

PROMPT = """This is a phone screenshot (LinkedIn post, Naukri listing, job group chat, WhatsApp image, etc.). It may contain HR / recruiter contact details.

Extract the contact details and return ONLY JSON, no markdown, in this shape:
{"entries":[{"company":"","profile":"","hr_name":"","emails":[],"phones":[]}]}

Rules:
- One entry per company / job post. If the image shows several posts, return several entries.
- emails: every email address for that post that is literally visible. Never guess or complete a partial email. If unreadable, skip it.
- Normalize obfuscated forms like "hr [at] abc [dot] com" to "hr@abc.com".
- phones: every mobile / WhatsApp / contact number visible for that post, written as shown (keep +91 or other country code if present). Never guess digits.
- company: the hiring company name if shown, else "".
- profile: the job title / role being hired for (e.g. ".NET Developer"), else "".
- hr_name: the recruiter / poster name if shown, else "".
- If the image has no email and no phone number, return {"entries":[]}."""


# ---------------------------------------------------------------- images / models
def encode_image(path: Path) -> str:
    img = Image.open(path).convert("RGB")
    img.thumbnail((2000, 2000))  # keeps small text readable, stays under size limits
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=88)
    return base64.standard_b64encode(buf.getvalue()).decode()


def make_chat(provider: str):
    """Returns chat(model, text, b64_image=None) -> str for the chosen provider."""
    if provider == "anthropic":
        import anthropic
        client = anthropic.Anthropic(max_retries=1)

        def chat(model, text, b64=None):
            content = []
            if b64:
                content.append({"type": "image", "source": {
                    "type": "base64", "media_type": "image/jpeg", "data": b64}})
            content.append({"type": "text", "text": text})
            msg = client.messages.create(
                model=model, max_tokens=1500,
                messages=[{"role": "user", "content": content}])
            return "".join(b.text for b in msg.content if b.type == "text")
        return chat

    from openai import OpenAI
    client = OpenAI(max_retries=1)

    def chat(model, text, b64=None):
        content = [{"type": "text", "text": text}]
        if b64:
            content.append({"type": "image_url",
                            "image_url": {"url": "data:image/jpeg;base64," + b64}})
        resp = client.chat.completions.create(
            model=model, max_tokens=1500,
            messages=[{"role": "user", "content": content}])
        return resp.choices[0].message.content or ""
    return chat


def parse_entries(text: str) -> list[dict]:
    m = re.search(r"\{.*\}", text, re.S)
    if not m:
        raise ValueError("model did not return JSON: " + text[:120])
    return json.loads(m.group(0)).get("entries", [])


# ---------------------------------------------------------------- cleaning helpers
def as_list(v) -> list[str]:
    if v is None:
        return []
    if isinstance(v, str):
        return [x.strip() for x in re.split(r"[,;\n]+", v) if x.strip()]
    return [str(x).strip() for x in v if str(x).strip()]


def clean_email(x: str) -> str:
    x = x.strip().lower().rstrip(".,;")
    return x if EMAIL_RE.match(x) else ""


def clean_phone(x: str) -> str:
    x = str(x).strip()
    digits = re.sub(r"\D", "", x)
    if not 8 <= len(digits) <= 15:
        return ""
    return ("+" if x.startswith("+") else "") + digits


def phone_key(p: str) -> str:
    return re.sub(r"\D", "", p)[-10:]  # same number with/without country code = same key


def company_key(name: str) -> str:
    n = re.sub(r"[^a-z0-9 ]", " ", name.lower())
    n = LEGAL_RE.sub(" ", n)  # "Acme Pvt. Ltd." and "ACME" are the same company
    return " ".join(n.split())


def add_unique(lst: list[str], value: str) -> None:
    v = (value or "").strip()
    if v and v.lower() not in {x.lower() for x in lst}:
        lst.append(v)


# ---------------------------------------------------------------- storage (one row per company)
class Store:
    def __init__(self):
        self.records = {}      # company key -> record
        self.emails = set()
        self.phones = set()    # phone keys

    def add(self, company, emails, phones, profiles, hr_names, sources) -> int:
        """Merge one entry. Returns number of NEW emails + phones stored."""
        company = (company or "").strip()
        new_emails, new_phones = [], []
        for e in map(clean_email, emails):
            if e and e not in self.emails and e not in new_emails:
                new_emails.append(e)
        for p in map(clean_phone, phones):
            if p and phone_key(p) not in self.phones and phone_key(p) not in map(phone_key, new_phones):
                new_phones.append(p)

        key = company_key(company)
        rec = self.records.get(key) if key else None
        if rec is None:
            if not (new_emails or new_phones):
                return 0
            rec_key = key or "nocompany:" + (new_emails[0] if new_emails else phone_key(new_phones[0]))
            rec = {"company": company, "emails": [], "phones": [],
                   "profiles": [], "hr_names": [], "sources": []}
            self.records[rec_key] = rec
        if company and not rec["company"]:
            rec["company"] = company

        for e in new_emails:
            rec["emails"].append(e)
            self.emails.add(e)
        for p in new_phones:
            rec["phones"].append(p)
            self.phones.add(phone_key(p))
        for lst, values in ((rec["profiles"], profiles), (rec["hr_names"], hr_names), (rec["sources"], sources)):
            for v in values:
                add_unique(lst, v)
        return len(new_emails) + len(new_phones)

    def rows(self):
        for r in self.records.values():
            yield [r["company"], TEXT_SEP.join(r["profiles"]), LIST_SEP.join(r["emails"]),
                   LIST_SEP.join(r["phones"]), TEXT_SEP.join(r["hr_names"]), TEXT_SEP.join(r["sources"])]


def split_text(v: str) -> list[str]:
    return [x.strip() for x in (v or "").split("|") if x.strip()]


def load_store(csv_path: Path) -> Store:
    """Loads an existing CSV (new format, or the older one-row-per-email format)."""
    store = Store()
    if csv_path.exists():
        with open(csv_path, newline="", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                store.add(
                    r.get("company"),
                    as_list(r.get("emails") or r.get("email")),
                    as_list(r.get("phones") or r.get("phone")),
                    split_text(r.get("profile")),
                    split_text(r.get("hr_name")),
                    split_text(r.get("source_file")),
                )
    return store


def write_csv(store: Store, path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(FIELDS)
        w.writerows(store.rows())


def write_excel(store: Store, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "HR Contacts"
    ws.append(FIELDS)
    for row in store.rows():
        ws.append(row)  # plain strings, so phone numbers stay text
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, width in zip("ABCDEF", (28, 30, 50, 28, 24, 30)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


# ---------------------------------------------------------------- main
def main(argv=None, chat_factory=make_chat) -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--provider", choices=["anthropic", "openai"], default="anthropic")
    ap.add_argument("--model", help="Model ID (defaults depend on provider)")
    ap.add_argument("--input", help="Folder containing screenshots")
    ap.add_argument("--output", default="hr_contacts", help="Output name (no extension)")
    ap.add_argument("--limit", type=int, help="Only process the first N new screenshots")
    ap.add_argument("--test", action="store_true", help="Send one text message to check the connection, then exit")
    args = ap.parse_args(argv)
    model = args.model or DEFAULT_MODELS[args.provider]

    load_dotenv(Path(__file__).with_name(".env"), override=True)
    load_dotenv(override=True)

    key_vars = ["ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN"] if args.provider == "anthropic" else ["OPENAI_API_KEY"]
    if not any(os.environ.get(k) for k in key_vars):
        sys.exit(f"{' or '.join(key_vars)} not found. Add it to your .env file.")
    base = os.environ.get("ANTHROPIC_BASE_URL" if args.provider == "anthropic" else "OPENAI_BASE_URL", "(provider default)")
    print(f"Provider: {args.provider} | model: {model} | base URL: {base}")

    chat = chat_factory(args.provider)

    if args.test:
        try:
            print("Reply:", chat(model, "Reply with the single word OK").strip()[:60])
            print("Connection works.")
        except Exception as e:
            sys.exit(f"Test failed: {e}")
        return

    if not args.input:
        sys.exit("--input is required (folder with screenshots)")
    folder = Path(args.input)
    if not folder.is_dir():
        sys.exit(f"Folder not found: {folder}")

    csv_path = Path(args.output + ".csv")
    xlsx_path = Path(args.output + ".xlsx")
    log_path = Path(args.output + "_processed.txt")
    locked = (f"Cannot write {csv_path} / {xlsx_path}. Close them in Excel, then run again "
              "(finished images are remembered, so nothing is lost).")

    store = load_store(csv_path)
    done = set(log_path.read_text(encoding="utf-8").splitlines()) if log_path.exists() else set()
    images = sorted(p for p in folder.iterdir() if p.suffix.lower() in EXTS and p.name not in done)
    if args.limit is not None and args.limit < 1:
        sys.exit("--limit must be greater than zero")
    if args.limit:
        images = images[:args.limit]
    print(f"{len(images)} screenshots to process "
          f"({len(store.records)} companies, {len(store.emails)} emails, {len(store.phones)} phones already saved)")

    consecutive_fails = 0
    try:
        write_csv(store, csv_path)  # makes sure the file exists / is writable
    except PermissionError:
        sys.exit(locked)

    with open(log_path, "a", encoding="utf-8") as log:
        for i, img in enumerate(images, 1):
            try:
                entries = parse_entries(chat(model, PROMPT, encode_image(img)))
                consecutive_fails = 0
            except Exception as e:
                consecutive_fails += 1
                print(f"[{i}/{len(images)}] {img.name}: FAILED ({str(e)[:200]})")
                if consecutive_fails >= 3:
                    sys.exit("\nStopping after 3 failures in a row - fix the error above and re-run "
                             "(finished images are remembered, so nothing is lost).")
                continue  # not logged as done, so it is retried next run

            added = 0
            for e in entries:
                added += store.add(
                    e.get("company"),
                    as_list(e.get("emails") or e.get("email")),
                    as_list(e.get("phones") or e.get("phone")),
                    [e.get("profile") or ""], [e.get("hr_name") or ""], [img.name],
                )
            try:
                write_csv(store, csv_path)
            except PermissionError:
                sys.exit(locked)
            log.write(img.name + "\n")
            log.flush()
            print(f"[{i}/{len(images)}] {img.name}: {added} new email/phone(s)")

    try:
        write_excel(store, xlsx_path)
    except PermissionError:
        sys.exit(locked)
    print(f"\nDone. {len(store.records)} companies, {len(store.emails)} emails, {len(store.phones)} phones. "
          f"Saved: {csv_path} and {xlsx_path}")


if __name__ == "__main__":
    main()